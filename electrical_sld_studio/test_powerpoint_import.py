import io
import json
import base64
from backend.powerpoint_exporter import export_network_to_powerpoint
from backend.powerpoint_importer import import_powerpoint_presentation, extract_text_and_tables_from_pptx
from backend.project_manager import load_project
from pptx import Presentation
from pptx.util import Inches

def run_tests():
    print("=== Testing PowerPoint Import & Auto-Layout ===")
    # 1. Test lossless roundtrip
    demo_proj = load_project("feeder_1789823077015")
    if not demo_proj:
        print("Demo project feeder_1789823077015 not found, skipping roundtrip test")
    else:
        metrics = {
            "total_length": 152.0, "ohl_length": 0, "ugc_length": 152.0,
            "total_capacity_kva": 210, "total_actual_load_kva": 150,
            "total_current_a": 11.0, "max_voltage_drop_pct": 0.8
        }
        pptx_bytes = export_network_to_powerpoint(demo_proj, metrics)
        assert len(pptx_bytes) > 0, "Failed to generate PPTX bytes"
        print(f"Generated PPTX size: {len(pptx_bytes)} bytes")

        # Import back
        stream = io.BytesIO(pptx_bytes)
        imported_proj, summary = import_powerpoint_presentation(stream)
        assert imported_proj is not None, "Failed to import PPTX"
        assert summary["mode"] == "lossless_embedded", f"Expected lossless_embedded mode, got {summary['mode']}"
        assert imported_proj["name"] == demo_proj["name"], "Project name mismatch"
        assert len(imported_proj["nodes"]) == len(demo_proj["nodes"]), "Nodes count mismatch"
        print(f"Roundtrip Test Passed! Imported {len(imported_proj['nodes'])} nodes, {len(imported_proj['sections'])} sections.")

    # 2. Test Generic PPTX with Table (No embedded metadata)
    prs = Presentation()
    prs.slide_width = Inches(8.27)
    prs.slide_height = Inches(11.69)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Add table with columns: [م, النود, اسم المحول / الكشك, النوع, القدرة, التحميل, الحمل]
    tbl_shape = slide.shapes.add_table(5, 7, Inches(0.5), Inches(1), Inches(7), Inches(5))
    tbl = tbl_shape.table
    headers = ["م", "النود", "اسم المحول / الكشك", "النوع", "القدرة", "التحميل", "الحمل"]
    for i, h in enumerate(headers):
        tbl.cell(0, i).text = h
    
    rows = [
        ("1", "N1", "محول المدرسة", "محول", "100", "70%", "70"),
        ("2", "N2", "كشك السوق الرئيسي", "كشك", "500", "60%", "300"),
        ("3", "N3", "محول العمدة", "محول", "160", "80%", "128"),
        ("4", "N4", "كشك السلام", "كشك", "300", "50%", "150")
    ]
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            tbl.cell(r_idx, c_idx).text = val

    out_stream = io.BytesIO()
    prs.save(out_stream)
    out_stream.seek(0)
    generic_bytes = out_stream.getvalue()

    # Import generic PPTX
    stream2 = io.BytesIO(generic_bytes)
    imported_gen, sum_gen = import_powerpoint_presentation(stream2)
    assert imported_gen is not None, "Failed to import generic PPTX"
    assert sum_gen["mode"] == "systematic_autolayout", f"Expected systematic_autolayout, got {sum_gen['mode']}"
    assert sum_gen["transformers_count"] == 2, f"Expected 2 transformers, got {sum_gen['transformers_count']}"
    assert sum_gen["kiosks_count"] == 2, f"Expected 2 kiosks, got {sum_gen['kiosks_count']}"
    assert sum_gen["switches_count"] >= 2, f"Expected switches created, got {sum_gen['switches_count']}"
    assert len(imported_gen["sections"]) > 0, "Expected sections created"
    
    # Verify auto-layout nodes have valid x, y
    for node in imported_gen["nodes"]:
        assert "x" in node and "y" in node, f"Node {node['id']} missing coordinates"
        assert node["x"] > 0 and node["y"] > 0, f"Invalid coords for node {node['id']}"

    print(f"Generic PPTX Test Passed! Extracted 2 transformers, 2 kiosks. Created {len(imported_gen['nodes'])} nodes, {len(imported_gen['sections'])} sections with systematic auto-layout.")
    print("ALL TESTS PASSED WITH 100% SUCCESS!")

if __name__ == "__main__":
    run_tests()
