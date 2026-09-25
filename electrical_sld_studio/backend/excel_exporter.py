"""
محرك تصدير تقارير الإكسيل الهندسية المتطورة
Copyright (C) ENG-MOSTAFAELMGHRABY - All Rights Reserved
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
ACCENT_FILL = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

FONT_HEADER = Font(name="Arial", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Arial", size=10, bold=True, color="000000")
FONT_NORMAL = Font(name="Arial", size=10, color="000000")
FONT_TITLE = Font(name="Arial", size=14, bold=True, color="1A365D")
FONT_COPYRIGHT = Font(name="Arial", size=9, italic=True, color="718096")

BORDER_THIN = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)
BORDER_DOUBLE_BOTTOM = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='double', color='1A365D')
)

def export_network_to_excel(project_data, metrics_data):
    """
    توليد مصنف إكسيل احترافي متكامل بمظهر هندسي راقٍ
    """
    wb = Workbook()
    
    # ورقة الأحمال Loads
    ws_loads = wb.active
    ws_loads.title = "Loads"
    ws_loads.views.sheetView[0].rightToLeft = True
    
    # عنوان الورقة
    ws_loads.merge_cells("A1:G1")
    title_cell = ws_loads["A1"]
    feeder_name = project_data.get("name", "مغذي التوزيع")
    title_cell.value = f"بيان أحمال ومحولات: {feeder_name}"
    title_cell.font = FONT_TITLE
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_loads.row_dimensions[1].height = 30
    
    # رؤوس الأعمدة
    headers_loads = [
        "رقم النود", "اسم المحول / الكشك", "النوع", "القدرة (KVA)", 
        "نسبة التحميل (%)", "الحمل الفعلي (KVA)", "تيار الجهد المتوسط (A)"
    ]
    
    for col_idx, h_text in enumerate(headers_loads, start=1):
        cell = ws_loads.cell(row=3, column=col_idx, value=h_text)
        cell.fill = HEADER_FILL
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN
    ws_loads.row_dimensions[3].height = 25
    
    # بيانات الأحمال
    loads_list = metrics_data.get("loads", [])
    row_idx = 4
    for ld in loads_list:
        ws_loads.cell(row=row_idx, column=1, value=ld["node_id"]).alignment = Alignment(horizontal="center")
        ws_loads.cell(row=row_idx, column=2, value=ld["name"]).alignment = Alignment(horizontal="right")
        ws_loads.cell(row=row_idx, column=3, value=ld["type"]).alignment = Alignment(horizontal="center")
        ws_loads.cell(row=row_idx, column=4, value=ld["capacity_kva"]).alignment = Alignment(horizontal="center")
        ws_loads.cell(row=row_idx, column=5, value=f"{ld['loading_pct']}%").alignment = Alignment(horizontal="center")
        ws_loads.cell(row=row_idx, column=6, value=ld["actual_load_kva"]).alignment = Alignment(horizontal="center")
        ws_loads.cell(row=row_idx, column=7, value=ld["actual_amp_mv"]).alignment = Alignment(horizontal="center")
        
        for col_idx in range(1, 8):
            c = ws_loads.cell(row=row_idx, column=col_idx)
            c.font = FONT_NORMAL
            c.border = BORDER_THIN
            if row_idx % 2 == 0:
                c.fill = ACCENT_FILL
        row_idx += 1
        
    # صف الإجمالي
    ws_loads.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
    tot_label = ws_loads.cell(row=row_idx, column=1, value="الإجمالـــــي")
    tot_label.font = FONT_BOLD
    tot_label.alignment = Alignment(horizontal="center", vertical="center")
    
    tot_cap = ws_loads.cell(row=row_idx, column=4, value=f"=SUM(D4:D{row_idx-1})")
    tot_cap.font = FONT_BOLD
    tot_cap.alignment = Alignment(horizontal="center")
    
    ws_loads.cell(row=row_idx, column=5, value="-").alignment = Alignment(horizontal="center")
    
    tot_act = ws_loads.cell(row=row_idx, column=6, value=f"=SUM(F4:F{row_idx-1})")
    tot_act.font = FONT_BOLD
    tot_act.alignment = Alignment(horizontal="center")
    
    tot_amp = ws_loads.cell(row=row_idx, column=7, value=f"=SUM(G4:G{row_idx-1})")
    tot_amp.font = FONT_BOLD
    tot_amp.alignment = Alignment(horizontal="center")
    
    for c_idx in range(1, 8):
        c = ws_loads.cell(row=row_idx, column=c_idx)
        c.fill = TOTAL_FILL
        c.border = BORDER_DOUBLE_BOTTOM
    ws_loads.row_dimensions[row_idx].height = 24
    
    # حقوق النشر في أسفل الشيت
    copy_row = row_idx + 2
    ws_loads.cell(row=copy_row, column=1, value="تم استخراج هذا التقرير آلياً بواسطة نظام Smart SLD Studio - جميع الحقوق محفوظة © ENG-MOSTAFAELMGHRABY").font = FONT_COPYRIGHT

    # ورقة المقاطع Sections
    ws_sec = wb.create_sheet(title="Sections")
    ws_sec.views.sheetView[0].rightToLeft = True
    
    ws_sec.merge_cells("A1:E1")
    title_sec = ws_sec["A1"]
    title_sec.value = f"بيان أطوال ومقاطع الخطوط والكابلات: {feeder_name}"
    title_sec.font = FONT_TITLE
    title_sec.alignment = Alignment(horizontal="center", vertical="center")
    ws_sec.row_dimensions[1].height = 30
    
    headers_sec = ["من نود", "إلى نود", "نوع الخط", "صنعة المقطع", "الطول (متر)"]
    for col_idx, h_text in enumerate(headers_sec, start=1):
        cell = ws_sec.cell(row=3, column=col_idx, value=h_text)
        cell.fill = HEADER_FILL
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN
    ws_sec.row_dimensions[3].height = 25
    
    sec_list = metrics_data.get("sections", [])
    row_idx = 4
    for sc in sec_list:
        ws_sec.cell(row=row_idx, column=1, value=sc["from_node"]).alignment = Alignment(horizontal="center")
        ws_sec.cell(row=row_idx, column=2, value=sc["to_node"]).alignment = Alignment(horizontal="center")
        ws_sec.cell(row=row_idx, column=3, value=sc["type"]).alignment = Alignment(horizontal="center")
        ws_sec.cell(row=row_idx, column=4, value=sc["size"]).alignment = Alignment(horizontal="center")
        ws_sec.cell(row=row_idx, column=5, value=sc["length"]).alignment = Alignment(horizontal="center")
        
        for col_idx in range(1, 6):
            c = ws_sec.cell(row=row_idx, column=col_idx)
            c.font = FONT_NORMAL
            c.border = BORDER_THIN
            if row_idx % 2 == 0:
                c.fill = ACCENT_FILL
        row_idx += 1
        
    # صف إجمالي الأطوال
    ws_sec.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
    tot_sec_lbl = ws_sec.cell(row=row_idx, column=1, value="إجمالي أطوال شبكة المغذي (متر)")
    tot_sec_lbl.font = FONT_BOLD
    tot_sec_lbl.alignment = Alignment(horizontal="center", vertical="center")
    
    tot_len = ws_sec.cell(row=row_idx, column=5, value=f"=SUM(E4:E{row_idx-1})")
    tot_len.font = FONT_BOLD
    tot_len.alignment = Alignment(horizontal="center")
    
    for c_idx in range(1, 6):
        c = ws_sec.cell(row=row_idx, column=c_idx)
        c.fill = TOTAL_FILL
        c.border = BORDER_DOUBLE_BOTTOM
    ws_sec.row_dimensions[row_idx].height = 24
    
    copy_row = row_idx + 2
    ws_sec.cell(row=copy_row, column=1, value="تم استخراج هذا التقرير آلياً بواسطة نظام Smart SLD Studio - جميع الحقوق محفوظة © ENG-MOSTAFAELMGHRABY").font = FONT_COPYRIGHT

    # ورقة الملخص الهندسي Summary
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].rightToLeft = True
    
    ws_sum.merge_cells("A1:B1")
    title_sum = ws_sum["A1"]
    title_sum.value = "الملخص الفني والهندسي للمغذي"
    title_sum.font = FONT_TITLE
    title_sum.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 30
    
    summary_data = [
        ("اسم المغذي (Feeder Name)", project_data.get("name", "مغذي رئيسي")),
        ("محطة التغذية / لوحة التوزيع", project_data.get("substation", "محطة محولات")),
        ("جهد التشغيل الاسمي", f"{project_data.get('voltage_kv', 11)} ك.ف"),
        ("إجمالي أطوال المغذي", f"{metrics_data['summary']['total_feeder_length_m']} متر ({metrics_data['summary']['total_feeder_length_m']/1000:.2f} كم)"),
        ("إجمالي أطوال الخطوط الهوائية", f"{metrics_data['summary']['total_ohl_length_m']} متر"),
        ("إجمالي أطوال الكابلات الأرضية", f"{metrics_data['summary']['total_ugc_length_m']} متر"),
        ("إجمالي سعة المحولات المركبة", f"{metrics_data['summary']['total_capacity_kva']} ك.ف.أ"),
        ("إجمالي الحمل الفعلي المتوقع", f"{metrics_data['summary']['total_actual_load_kva']} ك.ف.أ"),
        ("إجمالي تيار المغذي الفعلي", f"{metrics_data['summary']['total_feeder_current_a']} أمبير"),
        ("أقصى نسبة هبوط في الجهد", f"{metrics_data['summary']['max_voltage_drop_pct']}%"),
        ("مهندس التصميم والاعتماد", "ENG-MOSTAFAELMGHRABY"),
        ("حقوق الملكية الفكرية", "جميع الحقوق محفوظة للمهندس مصطفى المغربي © ENG-MOSTAFAELMGHRABY")
    ]
    
    for idx, (label, val) in enumerate(summary_data, start=3):
        c1 = ws_sum.cell(row=idx, column=1, value=label)
        c1.font = FONT_BOLD
        c1.fill = ACCENT_FILL
        c1.border = BORDER_THIN
        
        c2 = ws_sum.cell(row=idx, column=2, value=val)
        c2.font = FONT_NORMAL
        c2.border = BORDER_THIN
        ws_sum.row_dimensions[idx].height = 22

    # ضبط عروض الأعمدة تلقائياً في كل الأوراق
    for ws in [ws_loads, ws_sec, ws_sum]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len and len(val_str) < 50:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
